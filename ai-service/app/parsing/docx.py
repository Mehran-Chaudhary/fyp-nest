"""DOCX extraction on python-docx, walking the body in document order.

Word stores structure explicitly, so nothing is guessed: heading levels come
from the paragraph style (or its outline level), list items from numbering
properties, and tables are rendered as Markdown tables so the language model
reads rows and columns, not a run-on sentence. Content controls (`w:sdt`) are
descended into; headers, footers and comments are left out.
"""

from __future__ import annotations

import io
import re
import zipfile
from importlib.metadata import version
from typing import Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .. import errors
from .blocks import Block, ExtractedDocument, clean_text

PARSER_NAME = "python-docx"
_VERSION = version("python-docx")

OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_HEADING_STYLE = re.compile(r"^(?:heading|überschrift|titre|título|kop)\s*(\d)$", re.IGNORECASE)


def extract_docx(data: bytes) -> ExtractedDocument:
    if data.startswith(OLE_SIGNATURE):
        # Password-protected DOCX files are OLE containers holding an
        # encrypted package (legacy .doc files are too, but the backend has
        # already verified the type from the bytes).
        raise errors.encrypted()
    try:
        document = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError):
        raise errors.unparseable() from None

    blocks: list[Block] = []
    for element in _body_elements(document.element.body):
        if element.tag == qn("w:p"):
            block = _paragraph_block(Paragraph(element, document))
            if block is None:
                continue
            # Consecutive list items become one list block.
            if block.kind == "list" and blocks and blocks[-1].kind == "list":
                blocks[-1].text += "\n" + block.text
            else:
                blocks.append(block)
        elif element.tag == qn("w:tbl"):
            table = _table_markdown(Table(element, document))
            if table:
                blocks.append(Block("table", table))

    return ExtractedDocument(
        blocks=blocks,
        page_count=_saved_page_count(data),
        parser=PARSER_NAME,
        parser_version=_VERSION,
    )


def _body_elements(body) -> Iterator:  # type: ignore[no-untyped-def]
    for child in body.iterchildren():
        if child.tag in (qn("w:p"), qn("w:tbl")):
            yield child
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _body_elements(content)


def _paragraph_block(paragraph: Paragraph) -> Block | None:
    text = clean_text(paragraph.text)
    if not text:
        return None

    style = paragraph.style.name if paragraph.style is not None else ""
    level = _heading_level(paragraph, style)
    if level:
        return Block("heading", text, level=level)

    properties = paragraph._p.pPr  # noqa: SLF001 - python-docx exposes no public numbering API
    numbered = properties is not None and properties.numPr is not None
    if numbered or style.lower().startswith("list"):
        depth = 0
        if numbered and properties.numPr.ilvl is not None:
            depth = int(properties.numPr.ilvl.val or 0)
        return Block("list", "  " * depth + "- " + text)

    if style.lower() in ("quote", "intense quote"):
        return Block("quote", "> " + text)
    return Block("paragraph", text)


def _heading_level(paragraph: Paragraph, style: str) -> int:
    if style.lower() in ("title",):
        return 1
    match = _HEADING_STYLE.match(style.strip())
    if match:
        return max(1, min(int(match.group(1)), 6))
    properties = paragraph._p.pPr  # noqa: SLF001
    if properties is not None:
        outline = properties.find(qn("w:outlineLvl"))
        if outline is not None:
            value = outline.get(qn("w:val"))
            if value is not None and value.isdigit() and int(value) < 9:
                return min(int(value) + 1, 6)
    return 0


def _table_markdown(table: Table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            # Horizontally merged cells repeat the same element: keep one.
            if previous is not None and cell._tc is previous:  # noqa: SLF001
                continue
            previous = cell._tc  # noqa: SLF001
            text = clean_text(" ".join(p.text for p in cell.paragraphs)).replace("|", "\\|")
            cells.append(text.replace("\n", " "))
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join(lines)


def _saved_page_count(data: bytes) -> int | None:
    """The page count Word recorded at last save (docProps/app.xml), if any."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("docProps/app.xml").decode("utf-8", "ignore")
    except (KeyError, zipfile.BadZipFile):
        return None
    match = re.search(r"<Pages>(\d+)</Pages>", xml)
    return int(match.group(1)) if match and int(match.group(1)) > 0 else None
