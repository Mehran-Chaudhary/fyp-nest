"""Extraction: structure recovered from real PDF, DOCX, Markdown and text files."""

from __future__ import annotations

import io

import pytest

from app import errors
from app.parsing import extract
from app.parsing.text import decode_text

from .pdf_factory import TextRun, build_pdf, flow

BODY = 11.0


def kinds(document) -> list[tuple[str, str]]:  # type: ignore[no-untyped-def]
    return [(b.kind, b.text) for b in document.blocks]


# ── PDF ──────────────────────────────────────────────────────────────────────


def test_pdf_headings_paragraphs_and_lists() -> None:
    page = flow(
        [
            ("Employee Handbook", 20, True),
            ("", 0, False),
            ("1. Annual Leave", 14, True),
            ("", 0, False),
            ("Employees are entitled to twenty days of annual leave per", BODY, False),
            ("calendar year, accrued monthly from the date of joining.", BODY, False),
            ("", 0, False),
            ("Unused leave may be carried forward with approval.", BODY, False),
            ("", 0, False),
            ("- Submit requests two weeks in advance", BODY, False),
            ("- Emergencies are exempt", BODY, False),
        ]
    )
    document = extract(build_pdf([page]), "pdf")
    assert document.page_count == 1
    blocks = kinds(document)
    assert blocks[0] == ("heading", "Employee Handbook")
    assert blocks[1] == ("heading", "1. Annual Leave")
    assert document.blocks[0].level < document.blocks[1].level
    assert (
        "paragraph",
        "Employees are entitled to twenty days of annual leave per calendar year, accrued monthly from the date of joining.",
    ) in blocks
    assert ("paragraph", "Unused leave may be carried forward with approval.") in blocks
    assert ("list", "- Submit requests two weeks in advance\n- Emergencies are exempt") in blocks


def test_pdf_hyphenation_is_undone_and_pages_tracked() -> None:
    first = flow([("The reimburse-", BODY, False)], top=100)
    second = flow([("ment policy applies to all staff.", BODY, False)], top=800)
    document = extract(build_pdf([first, second]), "pdf")
    paragraph = document.blocks[0]
    assert paragraph.text == "The reimbursement policy applies to all staff."
    assert (paragraph.page, paragraph.last_page) == (1, 2)


def test_pdf_running_headers_and_page_numbers_are_removed() -> None:
    pages = []
    for number in range(1, 5):
        runs = [TextRun("ACME Corp — Confidential", 56, 820, 9)]
        runs += flow([(f"Body text of page {number} explains the rules in detail.", BODY, False)], top=760)
        runs.append(TextRun(f"Page {number} of 4", 270, 30, 9))
        pages.append(runs)
    document = extract(build_pdf(pages), "pdf")
    text = document.text
    assert "Confidential" not in text
    assert "Page 1 of 4" not in text
    assert text.count("Body text of page") == 4


def test_scanned_pdf_reports_empty_pages() -> None:
    document = extract(build_pdf([[], []]), "pdf")
    assert document.blocks == []
    assert document.empty_pages == 2


def test_not_a_pdf_is_unparseable() -> None:
    with pytest.raises(errors.ServiceError) as caught:
        extract(b"hello", "pdf")
    assert caught.value.status == 422 and caught.value.code == "UNPARSEABLE_DOCUMENT"


def test_corrupt_pdf_is_unparseable() -> None:
    with pytest.raises(errors.ServiceError) as caught:
        extract(b"%PDF-1.7\n" + b"\x00garbage" * 50, "pdf")
    assert caught.value.code == "UNPARSEABLE_DOCUMENT"


def test_password_protected_pdf_is_encrypted_document() -> None:
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter(clone_from=io.BytesIO(build_pdf([flow([("Secret", BODY, False)])])))
    writer.encrypt(user_password="open-sesame", owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    with pytest.raises(errors.ServiceError) as caught:
        extract(out.getvalue(), "pdf")
    assert caught.value.status == 422 and caught.value.code == "ENCRYPTED_DOCUMENT"


# ── DOCX ─────────────────────────────────────────────────────────────────────


def build_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees receive twenty days of annual leave.")
    document.add_heading("Sick Leave", level=2)
    document.add_paragraph("Submit a certificate after three days.", style="List Bullet")
    document.add_paragraph("Notify your manager the same day.", style="List Bullet")
    table = document.add_table(rows=3, cols=2)
    for row, values in enumerate([("Grade", "Days"), ("Staff", "20"), ("Manager | Senior", "25")]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_docx_structure() -> None:
    document = extract(build_docx(), "docx")
    blocks = kinds(document)
    assert blocks[0] == ("heading", "Leave Policy")
    assert document.blocks[0].level == 1
    assert ("heading", "Sick Leave") in blocks
    assert ("paragraph", "Employees receive twenty days of annual leave.") in blocks
    assert ("list", "- Submit a certificate after three days.\n- Notify your manager the same day.") in blocks
    table = next(text for kind, text in blocks if kind == "table")
    assert table.splitlines() == [
        "| Grade | Days |",
        "|---|---|",
        "| Staff | 20 |",
        "| Manager \\| Senior | 25 |",
    ]


def test_encrypted_docx_is_reported_as_such() -> None:
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
    with pytest.raises(errors.ServiceError) as caught:
        extract(ole, "docx")
    assert caught.value.code == "ENCRYPTED_DOCUMENT"


def test_broken_docx_is_unparseable() -> None:
    with pytest.raises(errors.ServiceError) as caught:
        extract(b"PK\x03\x04 not really a zip", "docx")
    assert caught.value.code == "UNPARSEABLE_DOCUMENT"


# ── Markdown and text ────────────────────────────────────────────────────────


MARKDOWN = """---
title: front matter is dropped
---
# Travel Policy

Economy class for flights under **six hours**.

## Per diem

| City | PKR |
|------|-----|
| Lahore | 5,000 |

```python
print("code stays whole")
```

- Book through the portal
- Keep receipts
"""


def test_markdown_structure() -> None:
    document = extract(MARKDOWN.encode(), "md")
    blocks = kinds(document)
    assert blocks[0] == ("heading", "Travel Policy")
    assert ("heading", "Per diem") in blocks
    assert ("paragraph", "Economy class for flights under **six hours**.") in blocks
    assert ("table", "| City | PKR |\n|------|-----|\n| Lahore | 5,000 |") in blocks
    assert any(kind == "code" and 'print("code stays whole")' in text for kind, text in blocks)
    assert ("list", "- Book through the portal\n- Keep receipts") in blocks
    assert "front matter" not in document.text


def test_plain_text_headings_and_paragraphs() -> None:
    text = "CODE OF CONDUCT\n\nBe respectful to colleagues.\nReport concerns early.\n\nConflicts of Interest\n\nDeclare them in writing."
    blocks = kinds(extract(text.encode("utf-8"), "txt"))
    assert blocks == [
        ("heading", "CODE OF CONDUCT"),
        ("paragraph", "Be respectful to colleagues.\nReport concerns early."),
        ("heading", "Conflicts of Interest"),
        ("paragraph", "Declare them in writing."),
    ]


@pytest.mark.parametrize(
    "data",
    [
        "Café — naïve".encode("utf-8"),
        "﻿Café — naïve".encode("utf-8"),
        "Café — naïve".encode("utf-16"),
        "Café — naïve".encode("utf-16-le"),
        "Café — naïve".encode("cp1252"),
    ],
)
def test_text_encodings_are_detected(data: bytes) -> None:
    assert decode_text(data).lstrip("﻿") == "Café — naïve"


def test_binary_garbage_is_not_text() -> None:
    with pytest.raises(errors.ServiceError):
        extract(bytes(range(256)) * 20, "txt")


def test_unsupported_type_is_415() -> None:
    with pytest.raises(errors.ServiceError) as caught:
        extract(b"x", "xlsx")
    assert caught.value.status == 415


def test_two_page_document_footer_is_removed() -> None:
    bodies = ["Annual leave is twenty days per year.", "Travel claims are paid within thirty days."]
    pages = [flow([(body, BODY, False)]) + [TextRun("Acme Handbook 2026", 56, 30, 9)] for body in bodies]
    text = extract(build_pdf(pages), "pdf").text
    assert "Acme Handbook 2026" not in text
    assert all(body in text for body in bodies)


def test_body_lines_outside_the_margins_are_never_treated_as_running() -> None:
    pages = [flow([("Definitions apply throughout this policy.", BODY, False)], top=500) for _ in range(4)]
    assert extract(build_pdf(pages), "pdf").text.count("Definitions apply") == 4


def test_scanned_pages_go_through_the_ocr_hook() -> None:
    from app.parsing.pdf import Line

    class FakeOcr:
        name = "fake"
        calls = 0

        def recognise_page(self, page, page_number):  # type: ignore[no-untyped-def]
            FakeOcr.calls += 1
            height = page.get_height()
            return [
                Line("Scanned Policy", 18.0, True, 56, 300, height - 60, height - 78, page_number),
                Line("Leave is twenty days per year.", 11.0, False, 56, 400, height - 100, height - 111, page_number),
            ]

    document = extract(build_pdf([[], [], []]), "pdf", ocr=FakeOcr(), max_ocr_pages=2)
    assert FakeOcr.calls == 2  # capped by max_ocr_pages
    assert document.ocr_pages == 2 and document.empty_pages == 1
    assert ("paragraph", "Leave is twenty days per year.") in kinds(document)
